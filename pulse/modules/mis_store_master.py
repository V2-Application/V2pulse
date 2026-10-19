"""MIS Store Master (DFS /sf-mis-store-master): the MIS team's own store master, AKS_SILVER.SLV_MIS_STORE_MASTER,
loaded from ST-MASTER.xlsx sheet ST. Port of DFS planning-service routes/akash/masterData.js (slug 'mis-store-master')
and db/aksPromote.js (MIS_STORE_MASTER): grid, distinct values, export, sample file, upload and delete-all.

Upload rules are DFS's: headers upper-cased (spaces / hyphens -> _), every file column must exist in the table,
ST_CD required, unique and upper-cased, typed cells checked (dates day-first: DD-MMM-YY(YY), DD-MM-YYYY,
YYYY-MM-DD; 'NA' and similar in a DATE / NUMBER cell = empty, a 0 date serial = empty), closed vocabularies for
ST_TYP / ST_STAT / ZONE. The file is appended to AKS_BRONZE.BRZ_MIS_STORE_MASTER (with LOAD_ID / LOAD_TS /
SRC_FILE / LOADED_BY when the table has them) and MERGEd into SILVER on ST_CD - a partial file never drops stores.
Delete-all clears SILVER and BRONZE (BRONZE is what a later promote reads)."""
import csv
import datetime as dt
import io
import re
import uuid

from fastapi import APIRouter, Depends, File, HTTPException, Request, UploadFile
from fastapi.responses import Response
from pydantic import BaseModel

from .. import audit
from ..iam.context import Ctx
from ..mis_engine import export as mx
from ..mis_engine.routes import sf_var_combined as base
from ..web import require

router = APIRouter(prefix="/api/v1/mis/store-master", tags=["mis"])
VIEW, EXPORT, ADMIN = "mis.view", "mis.export", "mis.admin"
SILVER = "V2RETAIL.AKS_SILVER.SLV_MIS_STORE_MASTER"
BRONZE = "V2RETAIL.AKS_BRONZE.BRZ_MIS_STORE_MASTER"
KEY = "ST_CD"
FILTER_COLS = ["ZONE", "REG", "AREA", "STATE", "RDC_CD", "ST_TYP", "ST_STAT", "OLD_NEW", "ZONE_RETAIL"]
SEARCH_COLS = ["ST_CD", "ST_NM", "ST_FULL_NM", "STATE"]
VALUES = {"ST_TYP": ["STORE", "NA", "DC", "HUB", "OFFICE-OLD", "OFFICE-NEW", "PLANT", "ZONAL_OFFICE"],
          "ST_STAT": ["ACT", "IN-ACT", "UPC", "NA", "SEDC", "SEST"],
          "ZONE": ["NORTH", "EAST", "NA"]}
NA_TOKENS = {"NA", "N/A", "#N/A", "NIL", "NONE", "-", "--"}
AUDIT_COLS = {"LOAD_ID", "LOAD_TS", "SRC_FILE", "LOADED_BY"}
IGNORABLE = {"UPDATED_AT", "CREATED_AT"} | AUDIT_COLS
NUMERIC = {"NUMBER", "DECIMAL", "NUMERIC", "INT", "INTEGER", "BIGINT", "SMALLINT", "TINYINT", "BYTEINT", "FLOAT",
           "FLOAT4", "FLOAT8", "DOUBLE", "REAL"}
MONTHS = {m: i for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}


def _cols(cur, fq):
    sch, tbl = fq.split(".")[1:]
    cur.execute("select COLUMN_NAME, DATA_TYPE from V2RETAIL.INFORMATION_SCHEMA.COLUMNS where TABLE_SCHEMA = %s and "
                "TABLE_NAME = %s order by ORDINAL_POSITION", (sch, tbl))
    return [(r[0].upper(), r[1].upper()) for r in cur.fetchall()]


# ── read ───────────────────────────────────────────────────────────────────────────────────────────
class GridReq(BaseModel):
    search: str | None = None
    col_filters: dict | None = None          # {COL: [values]}
    page: int = 1
    limit: int = 200
    sort_col: str | None = None
    sort_dir: str | None = "asc"


def _where(req: GridReq, cols):
    where, binds = [], []
    for c, vals in (req.col_filters or {}).items():
        if c in cols and isinstance(vals, list) and vals:
            vals = [str(v) for v in vals[:500]]
            blank = "(blank)" in vals
            vals = [v for v in vals if v != "(blank)"]
            parts = []
            if vals:
                parts.append("nvl(to_varchar(\"%s\"), '') in (%s)" % (c, ", ".join(["%s"] * len(vals))))
                binds += vals
            if blank:
                parts.append("nvl(to_varchar(\"%s\"), '') = ''" % c)
            where.append("(" + " or ".join(parts) + ")")
    s = (req.search or "").strip()
    if s:
        look = [c for c in SEARCH_COLS if c in cols]
        where.append("(" + " or ".join('upper(to_varchar("%s")) like %%s' % c for c in look) + ")")
        binds += ["%" + s.upper() + "%"] * len(look)
    sc = req.sort_col if req.sort_col in cols else KEY
    order = ' order by "%s" %s nulls last' % (sc, "desc" if (req.sort_dir or "").lower() == "desc" else "asc")
    return ((" where " + " and ".join(where)) if where else ""), binds, order


@router.get("/params")
def params(ctx: Ctx = Depends(require(VIEW))):
    conn = base._conn()
    try:
        cur = conn.cursor()
        cols = _cols(cur, SILVER)
        names = [c for c, _ in cols]
        opts = {}
        for c in [c for c in FILTER_COLS if c in names]:
            cur.execute("select distinct nvl(to_varchar(\"%s\"), '') from %s order by 1 limit 2000" % (c, SILVER))
            opts[c] = [r[0] if r[0] != "" else "(blank)" for r in cur.fetchall()]
        cur.execute(f"select count(*), max(UPDATED_AT) from {SILVER}" if "UPDATED_AT" in names else f"select count(*), null from {SILVER}")
        n, upd = cur.fetchone()
        return {"table": SILVER, "columns": names, "types": dict(cols), "filterCols": list(opts), "options": opts,
                "searchCols": SEARCH_COLS, "rows": n, "updatedAt": upd.isoformat() if upd else None,
                "values": VALUES, "key": KEY}
    finally:
        conn.close()


@router.post("/result")
def result(req: GridReq, ctx: Ctx = Depends(require(VIEW))):
    conn = base._conn()
    try:
        cur = conn.cursor()
        cols = [c for c, _ in _cols(cur, SILVER)]
        w, b, order = _where(req, cols)
        cur.execute(f"select count(*) from {SILVER}{w}", b)
        total = cur.fetchone()[0]
        limit, page = max(1, min(int(req.limit or 200), 5000)), max(1, int(req.page or 1))
        cur.execute(f"select * from {SILVER}{w}{order} limit {limit} offset {(page - 1) * limit}", b)
        return {"columns": cols, "rows": [[base._fmt_str(v) for v in r] for r in cur.fetchall()], "total": total,
                "page": page, "limit": limit}
    finally:
        conn.close()


@router.post("/export")
def export(req: GridReq, request: Request, fmt: str = "csv", ctx: Ctx = Depends(require(EXPORT))):
    conn = base._conn()
    try:
        cur = conn.cursor()
        cols = [c for c, _ in _cols(cur, SILVER)]
        w, b, order = _where(req, cols)
        cur.execute(f"select * from {SILVER}{w}{order}", b)
        if fmt == "xlsx":
            out = io.BytesIO()
            wb, fmts = mx.new_workbook(out)
            mx.add_sheet(wb, "MIS Store Master", cols, cur, fmts)
            wb.close()
            body, media, ext = out.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"
        else:
            out = io.StringIO()
            mx.write_csv(out, cols, cur)
            body, media, ext = out.getvalue(), "text/csv", "csv"
    finally:
        conn.close()
    audit.action_now(ctx, "mis.store_master.export", "mis", "store_master", {"format": fmt}, req=request)
    return Response(body, media_type=media, headers={"Content-Disposition": f'attachment; filename="MIS_STORE_MASTER.{ext}"'})


@router.get("/sample")
def sample(ctx: Ctx = Depends(require(VIEW))):
    """An empty upload file with the table's column headers (as DFS's -sample)."""
    conn = base._conn()
    try:
        cols = [c for c, _ in _cols(conn.cursor(), SILVER) if c not in IGNORABLE]
    finally:
        conn.close()
    out = io.BytesIO()
    wb, fmts = mx.new_workbook(out)
    ws = wb.add_worksheet("ST")
    for j, c in enumerate(cols):
        ws.write_string(0, j, c, fmts["head"])
    wb.close()
    return Response(out.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": 'attachment; filename="MIS_STORE_MASTER_SAMPLE.xlsx"'})


# ── upload ─────────────────────────────────────────────────────────────────────────────────────────
def _norm_header(s):
    return re.sub(r"[\[\]]", "", re.sub(r"[\s\-]+", "_", str(s).strip().upper()))


def _parse(data: bytes, name: str):
    """(headers, rows) of the first sheet (prefers a sheet named ST), or of a CSV."""
    if name.lower().endswith(".csv"):
        text = data.decode("utf-8-sig", errors="replace")
        rows = list(csv.reader(io.StringIO(text)))
    else:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb["ST"] if "ST" in wb.sheetnames else wb.worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    rows = [r for r in rows if any(v not in (None, "") for v in r)]
    if not rows:
        return [], []
    return [_norm_header(h) if h is not None else "" for h in rows[0]], rows[1:]


def _date(v):
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    s = str(v).strip()
    if re.match(r"^\d{4}-\d{2}-\d{2}", s):
        return s[:10]
    m = re.match(r"^(\d{1,2})[-/](\d{1,2})[-/](\d{2,4})$", s)
    if m:
        y = m.group(3)
        y = y if len(y) == 4 else ("19" + y if int(y) >= 50 else "20" + y)
        out = "%s-%02d-%02d" % (y, int(m.group(2)), int(m.group(1)))
    else:
        m = re.match(r"^(\d{1,2})[-\s]([A-Za-z]{3})[-\s](\d{2,4})$", s)
        if not m or m.group(2).lower() not in MONTHS:
            return None
        y = m.group(3)
        y = y if len(y) == 4 else ("19" + y if int(y) >= 50 else "20" + y)
        out = "%s-%02d-%02d" % (y, MONTHS[m.group(2).lower()], int(m.group(1)))
    try:
        d = dt.date.fromisoformat(out)
        return out if 1900 <= d.year <= 2100 else None
    except ValueError:
        return None


def _coerce(v, t, col, rowno, bad):
    if v is None or (isinstance(v, str) and v.strip() == ""):
        return None
    typed = t == "DATE" or t.startswith("TIMESTAMP") or t in NUMERIC
    if typed and str(v).strip().upper() in NA_TOKENS:
        return None
    if (t == "DATE" or t.startswith("TIMESTAMP")) and not isinstance(v, (dt.date, dt.datetime)):
        try:
            if float(str(v).strip()) == 0:
                return None
        except ValueError:
            pass
    if t == "DATE" or t.startswith("TIMESTAMP"):
        d = _date(v)
        if not d and len(bad) < 12:
            bad.append('row %d: %s = "%s" (expected a real date: DD-MMM-YYYY, DD-MM-YYYY or YYYY-MM-DD)' % (rowno, col, v))
        return d
    if t in NUMERIC:
        if isinstance(v, (int, float)):
            return v
        try:
            return float(str(v).strip().replace(",", ""))
        except ValueError:
            if len(bad) < 12:
                bad.append('row %d: %s = "%s" (expected a number)' % (rowno, col, v))
            return None
    return str(v).strip() if not isinstance(v, str) else v.strip()


@router.post("/upload")
async def upload(request: Request, file: UploadFile = File(...), ctx: Ctx = Depends(require(ADMIN))):
    data = await file.read()
    if len(data) > 25 * 1024 * 1024:
        raise HTTPException(413, "The file is larger than 25 MB.")
    from starlette.concurrency import run_in_threadpool
    out = await run_in_threadpool(_load, data, file.filename or "upload.xlsx", ctx.emp_code)
    audit.action_now(ctx, "mis.store_master.upload", "mis", "store_master", {"file": file.filename, **out}, req=request)
    return out


def _load(data, name, user):
    headers, rows = _parse(data, name)
    if not headers:
        raise HTTPException(400, "The file has no header row.")
    if not rows:
        raise HTTPException(400, "The file has a header but no data rows.")
    conn = base._conn()
    try:
        cur = conn.cursor()
        scols = dict(_cols(cur, SILVER))
        try:
            bcols = dict(_cols(cur, BRONZE))
        except Exception:  # noqa: BLE001
            bcols = {}
        if not bcols:
            raise HTTPException(503, f"V2Pulse cannot read {BRONZE} (no Snowflake grant). Nothing was loaded.")
        use = [(i, h) for i, h in enumerate(headers) if h in bcols and h in scols and h not in IGNORABLE]
        seen = set()
        use = [(i, h) for i, h in use if not (h in seen or seen.add(h))]
        unmatched = [h for h in headers if h and h not in bcols and h not in IGNORABLE]
        if unmatched:
            spare = [c for c in bcols if c not in [h for _, h in use] and c not in IGNORABLE]
            raise HTTPException(400, "Upload rejected - %d column(s) in the file do not exist in %s and would have loaded "
                                     "as blank: %s. %sRename the header(s) to match, or download the sample file."
                                % (len(unmatched), BRONZE, ", ".join('"%s"' % h for h in unmatched),
                                   ("Unused table column(s): %s. " % ", ".join(spare)) if spare else ""))
        if KEY not in [h for _, h in use]:
            raise HTTPException(400, f"The file is missing the key column {KEY}. Add it and re-upload.")
        bad, vals = [], []
        for rn, r in enumerate(rows, start=2):
            vals.append([_coerce(r[i] if i < len(r) else None, bcols[h], h, rn, bad) for i, h in use])
        if bad:
            raise HTTPException(400, "Upload rejected - cell(s) do not match the column type. Nothing was loaded. " + "; ".join(bad))
        ki = [h for _, h in use].index(KEY)
        seen_k, kbad = {}, []
        for rn, row in enumerate(vals, start=2):
            k = str(row[ki] or "").strip().upper()
            row[ki] = k or None
            if not k:
                kbad.append("row %d: %s is blank" % (rn, KEY))
            elif k in seen_k:
                kbad.append("row %d: %s = %s - already used on row %d" % (rn, KEY, k, seen_k[k]))
            else:
                seen_k[k] = rn
        if kbad:
            raise HTTPException(400, "Upload rejected - the file repeats or omits its key (%s). Nothing was loaded. %s"
                                % (KEY, "; ".join(kbad[:12])))
        vbad = []
        for col, allowed in VALUES.items():
            if col not in [h for _, h in use]:
                continue
            ci, ok = [h for _, h in use].index(col), {a.upper() for a in allowed}
            for rn, row in enumerate(vals, start=2):
                v = row[ci]
                if v not in (None, "") and str(v).strip().upper() not in ok and len(vbad) < 12:
                    vbad.append('row %d: %s = "%s" (allowed: %s)' % (rn, col, v, ", ".join(allowed)))
        if vbad:
            raise HTTPException(400, "Upload rejected - value(s) this master cannot hold. Nothing was loaded. " + "; ".join(vbad))
        load_id = uuid.uuid4().hex[:20]
        cols = [h for _, h in use]
        audit_cols = [c for c in ("LOAD_ID", "LOAD_TS", "SRC_FILE", "LOADED_BY") if c in bcols]
        audit_vals = {"LOAD_ID": load_id, "LOAD_TS": dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S"), "SRC_FILE": name,
                      "LOADED_BY": "V2Pulse:" + user}
        bsql = "insert into %s (%s) values (%s)" % (BRONZE, ", ".join('"%s"' % c for c in cols + audit_cols),
                                                     ", ".join(["%s"] * (len(cols) + len(audit_cols))))
        cur.execute("begin")
        try:
            for i in range(0, len(vals), 1000):
                cur.executemany(bsql, [row + [audit_vals[c] for c in audit_cols] for row in vals[i:i + 1000]])
            # MERGE the rows of THIS load into SILVER on ST_CD (only the columns the file carried are updated)
            src = "select %s from %s where LOAD_ID = %%s" % (", ".join('"%s"' % c for c in cols), BRONZE) if "LOAD_ID" in bcols \
                else None
            if src is None:
                raise HTTPException(503, f"{BRONZE} has no LOAD_ID column; cannot promote this load safely.")
            upd = ", ".join('t."%s" = s."%s"' % (c, c) for c in cols if c != KEY)
            stamp = ", t.UPDATED_AT = current_timestamp()::timestamp_ntz" if "UPDATED_AT" in scols else ""
            ins_cols = cols + (["UPDATED_AT"] if "UPDATED_AT" in scols else [])
            ins_vals = ['s."%s"' % c for c in cols] + (["current_timestamp()::timestamp_ntz"] if "UPDATED_AT" in scols else [])
            cur.execute(("merge into %s t using (%s) s on t.%s = s.%s when matched then update set %s%s "
                         "when not matched then insert (%s) values (%s)")
                        % (SILVER, src, KEY, KEY, upd, stamp, ", ".join('"%s"' % c for c in ins_cols), ", ".join(ins_vals)),
                        (load_id,))
            merged = cur.rowcount
            cur.execute(f"select count(*) from {SILVER}")
            total = cur.fetchone()[0]
            if total < 1:
                raise HTTPException(500, "After the load the master would be empty - rolled back.")
            cur.execute("commit")
        except Exception:
            cur.execute("rollback")
            raise
        return {"ok": True, "loadId": load_id, "rows": len(vals), "columns": cols, "merged": merged, "silverRows": total,
                "message": "Loaded %d store(s) from %s (%d column(s)); the master now holds %d stores." % (len(vals), name, len(cols), total)}
    finally:
        conn.close()


@router.delete("/all")
def delete_all(request: Request, ctx: Ctx = Depends(require(ADMIN))):
    """DFS 'Delete All Records': clears SILVER and BRONZE (BRONZE is what a later promote reads)."""
    conn = base._conn()
    try:
        cur = conn.cursor()
        cur.execute("begin")
        try:
            cur.execute(f"delete from {SILVER}")
            s = cur.rowcount
            cur.execute(f"delete from {BRONZE}")
            b = cur.rowcount
            cur.execute("commit")
        except Exception:
            cur.execute("rollback")
            raise
    finally:
        conn.close()
    audit.action_now(ctx, "mis.store_master.delete_all", "mis", "store_master", {"silver": s, "bronze": b}, req=request)
    return {"ok": True, "deleted": s, "message": "All MIS store master records deleted (%d)." % s}
