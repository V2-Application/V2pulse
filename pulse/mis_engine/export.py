"""MIS report output for V2Pulse: run a saved report into CSV / Excel, a multi-report workbook, or the
report's own Excel template. In DFS these lived in the Node planning-service (exceljs); the SQL comes
from the copied sf_mis_report exactly as its /export builds it, so the files match the grid."""
import csv
import datetime as _dt
import decimal
import io
import json
import os
import re

from .routes import sf_mis_report as rep
from .routes import sf_var_combined as base
from .routes.sf_gandola import _na_row, _text_cols

SHEET_MAX = 31
XLSX_MAX_ROWS = 1_048_575            # Excel's limit, less the header row
_BAD_SHEET = re.compile(r"[\[\]:*?/\\]")


def sheet_safe(name, used=None):
    """Excel sheet name rules (DFS lib/misBook.js): no []:*?/\\, at most 31 chars, unique."""
    s = _BAD_SHEET.sub("-", str(name or "Sheet")).strip().strip("'") or "Sheet"
    s = s[:SHEET_MAX]
    if used is not None:
        base_s, n = s, 2
        while s.upper() in used:
            tail = f" ({n})"
            s = base_s[:SHEET_MAX - len(tail)] + tail
            n += 1
        used.add(s.upper())
    return s


def file_safe(name):
    return re.sub(r"[^A-Za-z0-9_.() -]+", "_", str(name or "report")).strip() or "report"


def report_cursor(req: "rep.RunReq", runner=None):
    """Execute a saved (or unsaved) report the way /sf-mis-report/export does.
    Returns (conn, cur, header, meta); the caller fetches from `cur` and closes `conn`.
    `runner(conn, cur, sql)` executes the final statement (the Dump Scheduler passes one that can be stopped)."""
    conn = base._conn()
    cur = conn.cursor()
    try:
        meta = rep._resolve(cur, req)
        defn = meta["definition"]
        sql, keys, cols = rep.build_sql(defn, req.from_date, req.to_date, saved_conds=defn.get("col_filters"))
        where = rep._col_where_sql(cols, rep._outer_conds(defn, req.col_filters))
        turned = rep._is_turned(defn)
        sel = ", ".join(f'"{c}"' for c in cols) if turned else "*"
        tail = ', "_ORD"' if turned else ""
        order = (" ORDER BY " + ", ".join(keys) + tail) if keys else (' ORDER BY "_ORD"' if turned else "")
        final = f"SELECT {sel} FROM ({sql}){where}{order}"
        if runner:
            runner(conn, cur, final)
        else:
            cur.execute(final)
        hdr = [c[0] for c in cur.description]
        meta = {**meta, "keys": keys, "layout": rep._layout_of(defn)}
        return conn, cur, hdr, meta
    except Exception:
        conn.close()
        raise


def _rows(cur, batch=5000, tick=None):
    """Result rows with blank text cells as NA. `tick(n)` is called after each batch (progress / stop)."""
    text = _text_cols(cur)
    n = 0
    while True:
        if tick:
            tick(n)
        b = cur.fetchmany(batch)
        if not b:
            return
        n += len(b)
        for r in b:
            yield _na_row(r, text)


def _cell(v):
    if isinstance(v, decimal.Decimal):
        return float(v)
    return v


def write_csv(fh, hdr, cur, tick=None):
    """Same file as DFS /export (blank text cells as NA). Returns the row count."""
    w = csv.writer(fh)
    w.writerow(hdr)
    n = 0
    for r in _rows(cur, tick=tick):
        w.writerow(r)
        n += 1
    return n


def add_sheet(wb, name, hdr, cur, fmts, tick=None):
    """One report into one worksheet (constant-memory workbook). Returns the row count."""
    ws = wb.add_worksheet(name)
    for j, h in enumerate(hdr):
        ws.write_string(0, j, str(h), fmts["head"])
        ws.set_column(j, j, max(10, min(40, len(str(h)) + 2)))
    ws.freeze_panes(1, 0)
    n = 0
    for r in _rows(cur, tick=tick):
        n += 1
        if n > XLSX_MAX_ROWS:
            raise ValueError(f"'{name}' has more than {XLSX_MAX_ROWS:,} rows - too many for one Excel sheet. "
                             "Split it by a level or export CSV.")
        for j, v in enumerate(r):
            v = _cell(v)
            if v is None or v == "":
                continue
            if isinstance(v, (_dt.datetime, _dt.date)):
                ws.write_datetime(n, j, v if isinstance(v, _dt.datetime) else _dt.datetime(v.year, v.month, v.day), fmts["date"])
            elif isinstance(v, (int, float)) and not isinstance(v, bool):
                ws.write_number(n, j, v)
            else:
                ws.write_string(n, j, str(v))   # codes like plant 0001 stay text
    return n


def new_workbook(target):
    """`target`: a path or a BytesIO."""
    import xlsxwriter
    wb = xlsxwriter.Workbook(target, {"constant_memory": True, "strings_to_numbers": False,
                                      "default_date_format": "dd-mm-yyyy"})
    fmts = {"head": wb.add_format({"bold": True, "bg_color": "#FFF2CC", "border": 1}),
            "date": wb.add_format({"num_format": "dd-mm-yyyy"})}
    return wb, fmts


def write_xlsx(target, sheets, runner=None, tick=None):
    """`sheets`: list of (sheet name, RunReq). One workbook, one sheet per report. Returns
    [{sheet, name, rows}]."""
    wb, fmts = new_workbook(target)
    used, out = set(), []
    try:
        for sheet, req in sheets:
            conn, cur, hdr, meta = report_cursor(req, runner)
            try:
                name = sheet_safe(sheet or meta["name"], used)
                out.append({"sheet": name, "name": meta["name"], "rows": add_sheet(wb, name, hdr, cur, fmts, tick)})
            finally:
                conn.close()
    finally:
        wb.close()
    return out


# ── Excel templates (DFS: planning-service templates/MIS/<slug>.xlsx + .meta.json) ──────────────────
def _tpl_paths(folder, slug):
    p = os.path.join(folder, f"{slug}.xlsx")
    return p, p + ".meta.json"


def template_rows(folder, slug):
    try:
        with open(_tpl_paths(folder, slug)[1], encoding="utf-8") as fh:
            m = json.load(fh)
        if isinstance(m.get("nameRow"), int) and isinstance(m.get("dataRow"), int) and \
                (m["dataRow"] == 0 or m["dataRow"] > m["nameRow"]):
            return m
    except (OSError, ValueError):
        pass
    return {"nameRow": 1, "dataRow": 0}


def _row_names(ws, n):
    return [(c.column, str(c.value).strip().upper()) for c in ws[n] if c.value is not None and str(c.value).strip() != ""]


def best_header_row(ws, valid):
    """The row that names the most report columns (matches minus strangers, at least 3 matches)."""
    best = None
    for n in range(1, min(ws.max_row, 40) + 1):
        ns = [x for _, x in _row_names(ws, n)]
        if not ns:
            continue
        hit = sum(1 for x in ns if x in valid)
        score = hit - (len(ns) - hit)
        if hit >= 3 and (best is None or score > best["score"]):
            best = {"row": n, "hit": hit, "of": len(ns), "score": score}
    return best


def template_fill(tpl_path, meta, hdr, cur, max_rows=200_000):
    """Pour the report rows into its template: columns matched BY NAME on nameRow; dataRow 0 = append
    after the sheet's last row, otherwise overwrite from dataRow down copying that row's style."""
    import copy
    import openpyxl
    wb = openpyxl.load_workbook(tpl_path)
    ws = wb.worksheets[0]
    name_row, data_row = int(meta.get("nameRow") or 1), int(meta.get("dataRow") or 0)
    pos = {h.upper(): i for i, h in enumerate(hdr)}
    targets = [(col, pos[nm]) for col, nm in _row_names(ws, name_row) if nm in pos]
    start = data_row if data_row else ws.max_row + 1
    styles = {}
    if data_row:
        styles = {col: copy.copy(ws.cell(data_row, col)._style) for col, _ in targets}
    n = 0
    for r in _rows(cur):
        if n >= max_rows:
            raise ValueError(f"More than {max_rows:,} rows - too many for a formatted sheet. Split it or export CSV.")
        for col, i in targets:
            c = ws.cell(start + n, col)
            v = _cell(r[i])
            c.value = None if v == "" else v
            if col in styles:
                c._style = copy.copy(styles[col])
        n += 1
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), n
